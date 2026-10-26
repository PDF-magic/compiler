"""Regression checks for court and Federal Register Word exports."""
from io import BytesIO
from zipfile import ZipFile
import unittest
from docx import Document
from app import app
from word_export import build_word_document


class WordExportTests(unittest.TestCase):
    def test_court_editable_text_and_pages(self):
        payload = build_word_document(
            "# UNITED STATES DISTRICT COURT\n\nCase No. 24-1\n\n"
            "## Motion\n\n_See_ **authority**.[^a]\n\n[^a]: Citation note.",
            profile="court_docx", title="Motion", author="Party")
        doc = Document(BytesIO(payload))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("UNITED STATES DISTRICT COURT", text)
        self.assertIn("1. Citation note.", text)
        self.assertEqual(doc.core_properties.title, "Motion")
        self.assertEqual(doc.styles["Normal"].font.size.pt, 14)
        self.assertEqual(doc.sections[0].left_margin.inches, 1)
        with ZipFile(BytesIO(payload)) as zipfile:
            self.assertIn('w:instr="PAGE"', zipfile.read("word/footer1.xml").decode())

    def test_federal_register_safe_editable_draft(self):
        payload = build_word_document(
            "# DEPARTMENT OF EXAMPLE\n\nAGENCY: Example Office.\n\n"
            "ACTION: Notice.\n\nSUMMARY: [support](https://example.org/details).\n\n"
            "1. Literal instruction.\n\n"
            "SUPPLEMENTARY INFORMATION: Draft.[^n]\n\n[^n]: Manual note.",
            profile="federal_register_docx", title="Secret", author="Private")
        doc = Document(BytesIO(payload))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("AGENCY: Example Office.", text)
        self.assertIn("1. Literal instruction.", text)
        self.assertIn("support (https://example.org/details)", text)
        self.assertIn("1. Manual note.", text)
        self.assertEqual(doc.core_properties.title, "")
        self.assertEqual(doc.core_properties.author, "")
        self.assertEqual(doc.styles["Normal"].font.size.pt, 12)
        with ZipFile(BytesIO(payload)) as zipfile:
            xml = zipfile.read("word/document.xml")
            self.assertNotIn(b"<w:hyperlink", xml)
            self.assertNotIn(b"<w:fldSimple", xml)
            self.assertNotIn(b"<w:numPr", xml)

    def test_web_routes(self):
        with app.test_client() as client:
            word = client.post("/render", data={
                "markdown": "# DEPARTMENT\n\nAGENCY: Example.\n\nACTION: Notice.",
                "output_format": "federal_register_docx", "output_name": "notice.pdf"})
            self.assertEqual(word.status_code, 200)
            self.assertIn("wordprocessingml.document", word.mimetype)
            self.assertIn("notice.docx", word.headers["Content-Disposition"])
            self.assertTrue(word.data.startswith(b"PK"))
            pdf = client.post("/render", data={"markdown": "# Heading\n\nText."})
            self.assertEqual(pdf.status_code, 200)
            self.assertEqual(pdf.mimetype, "application/pdf")
            self.assertTrue(pdf.data.startswith(b"%PDF"))
            invalid = client.post("/render", data={
                "markdown": "Text", "output_format": "unsupported"})
            self.assertEqual(invalid.status_code, 400)


if __name__ == "__main__":
    unittest.main()
