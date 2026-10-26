"""Editable Word exports for court and Federal Register drafting profiles."""

from io import BytesIO
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

WORD_PROFILES = ("court_docx", "federal_register_docx")
CAPTIONS = ("AGENCY:", "ACTION:", "SUMMARY:", "DATES:", "ADDRESSES:",
            "FOR FURTHER INFORMATION CONTACT:", "SUPPLEMENTARY INFORMATION:")
NOTE = re.compile(r"^\[\^([A-Za-z0-9_-]+)\]:\s*(.*)$")
REF = re.compile(r"\[\^([A-Za-z0-9_-]+)\]")
INLINE = re.compile(
    r"\[\^[A-Za-z0-9_-]+\]|\[([^\]\n]+)\]\((https?://[^)\s]+)\)"
    r"|\*\*[^*\n]+\*\*|__[^_\n]+__|(?<!\w)_[^_\n]+_"
    r"|(?<!\*)\*[^*\n]+\*|\x60[^\x60\n]+\x60"
)
HEADING = re.compile(r"^(#{1,6})[ \t]+(.+)$")
LIST = re.compile(r"^[ \t]*(?:(\d+)[.)]|([-+*]))[ \t]+(.+)$")
FENCE = re.compile(r"^[ \t]*(\x60{3,}|~{3,})")


def _split_notes(text):
    lines, notes, active = [], {}, None
    for line in text.splitlines():
        match = NOTE.match(line)
        if match:
            active = match.group(1)
            notes[active] = match.group(2).strip()
        elif active and (line.startswith("    ") or line.startswith("\t")):
            notes[active] += " " + line.strip()
        else:
            active = None
            lines.append(line)
    return lines, notes


def _runs(paragraph, text, numbers):
    """Create editable runs; never create a live hyperlink relationship."""
    text = re.sub(
        r"\{\{([A-Za-z0-9_-]+)\}\}",
        lambda m: str(numbers[m.group(1)]) if m.group(1) in numbers else m.group(0),
        text,
    )
    offset = 0
    for match in INLINE.finditer(text):
        if match.start() > offset:
            paragraph.add_run(text[offset:match.start()])
        token = match.group(0)
        ref = REF.fullmatch(token)
        if ref and ref.group(1) in numbers:
            paragraph.add_run(str(numbers[ref.group(1)])).font.superscript = True
        elif match.group(1) is not None:
            paragraph.add_run(match.group(1) + " (" + match.group(2) + ")")
        elif token.startswith(("**", "__")):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith(("*", "_")):
            paragraph.add_run(token[1:-1]).italic = True
        elif token.startswith(chr(96)):
            paragraph.add_run(token[1:-1]).font.name = "Courier New"
        else:
            paragraph.add_run(token)
        offset = match.end()
    if offset < len(text):
        paragraph.add_run(text[offset:])


def build_word_document(markdown: str, *, profile: str,
                        title: str = "", author: str = "") -> bytes:
    """Build editable DOCX. Notes become numbered paragraphs, not Word footnotes.

    These are draft templates: local court rules and OFR manuscript requirements
    must be reviewed before any real filing or agency publication.
    """
    if profile not in WORD_PROFILES:
        raise ValueError("Unknown Word export profile.")
    if not markdown.strip():
        raise ValueError("Upload a Markdown file or paste Markdown text.")
    federal = profile == "federal_register_docx"
    lines, notes = _split_notes(markdown)
    numbers = {}
    in_fence, fence_char = False, ""
    for line in lines:
        marker = FENCE.match(line)
        if marker:
            if not in_fence:
                in_fence, fence_char = True, marker.group(1)[0]
            elif marker.group(1)[0] == fence_char:
                in_fence = False
            continue
        if not in_fence:
            for key in REF.findall(line):
                if key in notes and key not in numbers:
                    numbers[key] = len(numbers) + 1

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12 if federal else 14)
    normal.paragraph_format.line_spacing = 2
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    for level in range(1, 7):
        style = doc.styles["Heading " + str(level)]
        style.font.name = "Times New Roman"
        style.font.size = normal.font.size
        style.font.bold = True
        style.paragraph_format.line_spacing = 2
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.space_before = Pt(0)

    props = doc.core_properties
    if federal:
        # Avoid carrying private author/title into an OFR agency draft.
        for name in ("title", "subject", "author", "last_modified_by",
                     "keywords", "comments", "category", "identifier"):
            setattr(props, name, "")
    else:
        props.title, props.author = title, author
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)

    def add(text, kind="body", level=1):
        if not text:
            return
        para = doc.add_paragraph(style=("Heading " + str(min(level, 6)))
                                 if kind == "heading" else None)
        if kind == "heading" and level == 1 and not federal:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if kind in ("quote", "code"):
            para.paragraph_format.line_spacing = 1
        if kind == "quote":
            para.paragraph_format.left_indent = Inches(0.5)
            para.paragraph_format.right_indent = Inches(0.5)
        if kind == "list":
            para.paragraph_format.left_indent = Inches(0.25)
        if kind == "code":
            para.add_run(text).font.name = "Courier New"
        elif federal and kind == "body" and any(text.upper().startswith(k) for k in CAPTIONS):
            label = next(k for k in CAPTIONS if text.upper().startswith(k))
            para.add_run(label).bold = True
            _runs(para, text[len(label):], numbers)
        else:
            _runs(para, text, numbers)

    prose, quotes, code = [], [], []
    in_fence, fence_char = False, ""

    def flush(items, kind):
        if items:
            add(" ".join(items).strip(), kind)
            items.clear()

    for line in lines:
        marker = FENCE.match(line)
        if marker:
            flush(prose, "body")
            flush(quotes, "quote")
            if not in_fence:
                in_fence, fence_char = True, marker.group(1)[0]
            elif marker.group(1)[0] == fence_char:
                if code:
                    add("\n".join(code), "code")
                code.clear()
                in_fence = False
            continue
        if in_fence:
            code.append(line)
            continue
        stripped = line.strip()
        if not stripped:
            flush(prose, "body")
            flush(quotes, "quote")
            continue
        heading = HEADING.match(stripped)
        listed = LIST.match(line)
        if heading:
            flush(prose, "body")
            flush(quotes, "quote")
            add(heading.group(2), "heading", len(heading.group(1)))
        elif stripped == "[[TOC]]":
            flush(prose, "body")
            flush(quotes, "quote")
            # No active Word TOC fields, especially for OFR documents.
        elif stripped == "[[signature]]":
            flush(prose, "body")
            flush(quotes, "quote")
            if not federal:
                add("_______________________________")
        elif stripped.startswith(">"):
            flush(prose, "body")
            content = stripped[1:].lstrip()
            if content.startswith("[!") and "]" in content:
                if content.upper().startswith(("[!SOURCE]", "[!SRC]")):
                    flush(quotes, "quote")
                    add("— " + content.split("]", 1)[1].strip())
            else:
                quotes.append(content)
        elif listed:
            flush(prose, "body")
            flush(quotes, "quote")
            prefix = listed.group(1) + "." if listed.group(1) else "•"
            add(prefix + " " + listed.group(3), "list")
        else:
            flush(quotes, "quote")
            prose.append(stripped)
    flush(prose, "body")
    flush(quotes, "quote")
    if code:
        add("\n".join(code), "code")
    if numbers:
        add("Notes", "heading")
        for key, number in numbers.items():
            add(str(number) + ". " + notes[key])

    destination = BytesIO()
    doc.save(destination)
    return destination.getvalue()
